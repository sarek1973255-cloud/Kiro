"""
ExcelParser — reads a DELIXI-style customs Excel workbook and returns a
DeclarationData object.

Supported layouts
─────────────────
Layout A (two-sheet workbook):
  Sheet 'пакинг лист'  — Packing List  (weights, package count, totals)
  Sheet 'Лист1'        — Invoice        (item numbers, HS codes, prices)

Layout B (single-sheet workbook, "Invoys" / "Invoys (2)"):
  One sheet contains both Commercial Invoice (left half, cols 0-6) and
  Packing List (right half, cols 7+) side by side.  The same sheet is used
  for both "packing" and "invoice" parsing roles.

  The Всего totals row contains two 'Всего' labels: the first at col 0 is
  the invoice total, the second further right is the packing total followed
  by packages / gross-weight / net-weight.

The parser uses fuzzy sheet-name matching so minor spelling or case variations
are handled automatically.

BUG 2 fix (quantity becomes zero)
──────────────────────────────────
Root cause (manba.xlsx / DBT3004): The data rows have quantity and unit-measure
columns SWAPPED relative to the header row (col 3 has the numeric quantity, col 4
has the unit abbreviation 'шт').  The header-based column detector correctly maps
the quantity column to col 4 ('Кол-во / Q'ty'), but the actual numeric value is
in col 3 — so the parser reads 'шт' for quantity, which converts to 0.

Fixes applied:
  1. Expanded quantity column hints to cover more header variants.
  2. Apostrophe variants in "Q'ty" normalised to plain ASCII in _norm().
  3. After column mapping, _post_validate_columns() checks a sample of data
     rows; if the mapped quantity column contains non-numeric values while an
     adjacent unit column contains numbers, the two are swapped.
  4. If quantity is still 0 after the swap, it is derived from
     total_price / unit_price when both are non-zero.
  5. All silent-zero conversions emit a WARNING log entry.

BUG 1 fix (HS code duplication in product name)
  Fixed in translation.py — fallback returns "商品" without the HS code.

COMPANY NAME FIX
────────────────
  Company names read from the Excel often contain legal-form prefixes (OOO,
  ООО), surrounding guillemet quotes (« »), and tax-ID / registration numbers
  inside parentheses.  _clean_company_name() normalises these into a clean
  "Name LLC" form expected on the customs declaration.

  Examples:
    OOO «DEVELOPER BUSSINES TRASS» (309387445)  →  DEVELOPER BUSSINES TRASS LLC
    ООО "Elektro Fayz Plyus" (12345)             →  Elektro Fayz Plyus LLC
    Elektro Fayz Plyus LLC                        →  Elektro Fayz Plyus LLC  (unchanged)
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from typing import Optional

import openpyxl

from translation import get_chinese_name, normalize_hs

log = logging.getLogger(__name__)


# ── Company name normalisation ────────────────────────────────────────────

def _clean_company_name(name: str) -> str:
    """
    Normalise a company name extracted from the Excel / consignee field.

    Rules (per spec):
      1. Remove OOO / ООО prefix variants (both Cyrillic and Latin letter O).
         These indicate a Russian/CIS LLC legal form — the suffix "LLC" is added.
      2. Remove surrounding quotation marks (« », ", ', ").
      3. Remove tax IDs and registration numbers — any content inside
         parentheses, square brackets, or curly braces.
      4. Remove leading/trailing LLC prefix if it would be duplicated.
      5. Normalise runs of whitespace to a single space.
      6. Append " LLC" if the original name had an OOO/ООО prefix and the
         cleaned name does not already end with LLC / Ltd / Limited.

    Examples:
      "OOO «DEVELOPER BUSSINES TRASS» (309387445)" → "DEVELOPER BUSSINES TRASS LLC"
      "ООО Elektro Fayz Plyus (12345678)"          → "Elektro Fayz Plyus LLC"
      "Elektro Fayz Plyus LLC"                     → "Elektro Fayz Plyus LLC"
    """
    if not name:
        return name

    name = name.strip()
    had_ooo = False

    # Step 1: detect and remove OOO / ООО prefix.
    # The letters may be Cyrillic О (U+041E) or Latin O (U+004F) in any mix.
    _OOO_PATTERN = re.compile(
        r"^[ОO][ОO][ОO]\s*",
        re.IGNORECASE,
    )
    if _OOO_PATTERN.match(name):
        name = _OOO_PATTERN.sub("", name, count=1).strip()
        had_ooo = True

    # Step 2: remove parenthesised / bracketed tax IDs and registration numbers
    # BEFORE removing surrounding quotes, because the closing quote may appear
    # before the bracket: «NAME» (123456) → after brackets: «NAME» → then quotes.
    # e.g. "(309387445)", "[12345]", "{...}"
    name = re.sub(r'\s*[\(\[\{][^\)\]\}]*[\)\]\}]', "", name)
    name = name.strip()

    # Step 3: remove ALL guillemet and ASCII quote characters wherever they
    # appear in the name.  This handles cases like «NAME» LLC where the
    # closing » is not at the end of the string (e.g. «ASL FOOD» LLC (id)
    # becomes ASL FOOD LLC after bracket removal in step 2 leaves «ASL FOOD» LLC).
    name = re.sub(r'[«»\"\u201C\u201D\u2018\u2019\']+', '', name)
    name = name.strip()

    # Step 4: remove LLC / Ltd / Limited prefix if already at the start
    # (it will be appended at the end instead to ensure consistent placement).
    _LLC_PREFIX = re.compile(r"^(LLC|Ltd\.?|Limited)\s+", re.IGNORECASE)
    if _LLC_PREFIX.match(name):
        name = _LLC_PREFIX.sub("", name, count=1).strip()

    # Step 5: normalise whitespace
    name = re.sub(r"\s+", " ", name).strip()

    # Step 6: ensure LLC suffix when company had an OOO/ООО legal form
    _LLC_SUFFIX = re.compile(r"\b(LLC|Ltd\.?|Limited)\s*$", re.IGNORECASE)
    if had_ooo and not _LLC_SUFFIX.search(name):
        name = name + " LLC"

    return name


# ── Data models ──────────────────────────────────────────────────────────

@dataclass
class ProductItem:
    item_no:      int
    hs_code:      str
    russian_name: str
    chinese_name: str
    quantity:     int
    unit_price:   float
    total_price:  float
    unit:         str = "шт"    # Unit of measure from Excel (шт, kg/кг, etc.)

    @property
    def qty_str(self) -> str:
        """Format quantity with Chinese unit suffix."""
        unit_lower = self.unit.lower().strip()
        if unit_lower in ("kg", "кг", "kg/кг", "килограмм", "килограммы"):
            return f"{self.quantity}千克"
        elif unit_lower in ("m", "м", "метр", "метры"):
            return f"{self.quantity}米"
        elif unit_lower in ("set", "комплект", "компл", "компл."):
            return f"{self.quantity}套"
        else:
            return f"{self.quantity}个"

    @property
    def unit_price_str(self) -> str:
        return f"{self.unit_price:.4f}".rstrip("0").rstrip(".")

    @property
    def total_price_str(self) -> str:
        return f"{self.total_price:.2f}"

    @property
    def hs_name_str(self) -> str:
        """Legacy property kept for backward compatibility."""
        return f"{self.hs_code} {self.chinese_name}"


@dataclass
class DeclarationData:
    # Header fields
    contract_no:       str   = ""
    invoice_date:      str   = ""
    sender_name:       str   = ""
    sender_address:    str   = ""
    consignee:         str   = ""
    consignee_address: str   = ""
    delivery_terms:    str   = ""

    # Totals from Packing List
    packages:          int   = 0
    net_weight:        float = 0.0
    gross_weight:      float = 0.0

    # Fields filled by the caller (not present in Excel)
    declaration_date:  str   = ""   # YYYYMMDD
    pre_entry_no:      str   = ""   # 18-digit customs number
    export_date:       str   = ""

    # Product lines
    items: list[ProductItem] = field(default_factory=list)

    @property
    def total_amount_usd(self) -> float:
        return sum(i.total_price for i in self.items)


# ── Column-detection keywords ─────────────────────────────────────────────

_INVOICE_COL_HINTS: dict[str, list[str]] = {
    "item_no":     ["№", "item no", "item#", "item num", "порядков"],
    "description": ["description of", "наименование", "товар", "goods"],
    "hs_code":     ["hc code", "hs code", "тн вэд", "тнвэд", "hscode", "tnved"],
    "unit":        ["msr.unit", "ед. изм", "ед.изм", "measure"],
    # BUG 2: expanded quantity hints; apostrophe normalised in _norm()
    "quantity":    [
        "q'ty", "qty", "кол-во", "quantity", "количество",
        "кол", "шт", "штук", "pieces", "pcs", "число", "единиц",
    ],
    "unit_price":  ["unit price", "unit   price", "цена за", "price (usd)", "unit cost"],
    "total":       ["amount (usd)", "amount(usd)", "total", "сумма", "итого", "total price"],
}

_PACKING_COL_HINTS: dict[str, list[str]] = {
    "item_no":     ["№", "item no", "item#", "порядков"],
    "description": ["description of", "наименование", "товар", "goods"],
    "hs_code":     ["hc code", "hs code", "тн вэд", "тнвэд"],
    "quantity":    ["quantity", "кол-во", "qty", "кол", "шт", "штук", "pieces", "pcs"],
    "packages":    ["place", "мест", "pkg", "pcs"],
    "net_weight":  ["net weight", "нетто", "net w"],
    "gross_weight":["gross weight", "брутто", "gross w"],
}

# Default (fallback) 0-based column indices for the original format
_INVOICE_DEFAULT_COLS = {
    "item_no": 0, "description": 1, "hs_code": 2,
    "unit": 3, "quantity": 4, "unit_price": 5, "total": 6,
}
_PACKING_DEFAULT_COLS = {
    "item_no": 0, "description": 1, "hs_code": 2,
    "quantity": 3, "packages": 4, "net_weight": 5, "gross_weight": 6,
}


# ── Parser ────────────────────────────────────────────────────────────────

class ExcelParser:
    """Parse a DELIXI-style customs Excel workbook."""

    PACKING_HINTS = ["пакинг", "packing", "pack"]
    INVOICE_HINTS = ["лист1", "лист 1", "invoice", "sheet1", "инвойс", "invoys"]

    TOTALS_KEYWORD = "всего"

    def parse(self, path: str) -> DeclarationData:
        wb = openpyxl.load_workbook(path, data_only=True)
        packing_ws, invoice_ws = self._find_sheets(wb)

        data = DeclarationData()
        self._parse_header_rows(packing_ws, data)
        packing_header_row, pack_col_map = self._detect_columns(
            packing_ws, _PACKING_COL_HINTS, _PACKING_DEFAULT_COLS,
        )
        invoice_header_row, inv_col_map = self._detect_columns(
            invoice_ws, _INVOICE_COL_HINTS, _INVOICE_DEFAULT_COLS,
        )

        # BUG 2: Validate and potentially swap quantity/unit columns
        inv_col_map = self._post_validate_columns(
            invoice_ws, inv_col_map, invoice_header_row + 1,
        )

        log.info(
            "Column map — packing: %s  invoice: %s",
            pack_col_map, inv_col_map,
        )
        self._parse_totals(packing_ws, pack_col_map)
        self._fill_totals_from_sheet(packing_ws, data, pack_col_map)
        self._parse_items(invoice_ws, data, inv_col_map, invoice_header_row + 1)
        return data

    # ── Sheet resolution ───────────────────────────────────────────────────

    def _find_sheets(self, wb: openpyxl.Workbook):
        """
        Return (packing_sheet, invoice_sheet) using fuzzy name matching.

        Single-sheet workbooks (Layout B): If only one sheet is found, it is
        used for both packing and invoice roles.  Both halves of the combined
        sheet are parsed by their respective methods.
        """
        names = wb.sheetnames

        if len(names) == 1:
            log.info(
                "Single-sheet workbook detected (%r) — using same sheet "
                "for both packing and invoice roles.", names[0]
            )
            return wb[names[0]], wb[names[0]]

        packing = self._match_sheet(names, self.PACKING_HINTS, "packing list")
        invoice = self._match_sheet(names, self.INVOICE_HINTS, "invoice",
                                    exclude=packing)
        log.info("Sheets resolved — packing: %r  invoice: %r", packing, invoice)
        return wb[packing], wb[invoice]

    @staticmethod
    def _match_sheet(
        names: list[str],
        hints: list[str],
        role: str,
        exclude: Optional[str] = None,
    ) -> str:
        normed = {n: _norm(n) for n in names}
        for hint in hints:
            for n, nn in normed.items():
                if hint in nn and n != exclude:
                    return n
        for n in names:
            if n != exclude:
                return n
        raise ValueError(
            f"Could not find {role} sheet. Available sheets: {names}"
        )

    # ── Column detection ───────────────────────────────────────────────────

    @staticmethod
    def _detect_columns(
        ws,
        hints: dict[str, list[str]],
        defaults: dict[str, int],
    ) -> tuple[int, dict[str, int]]:
        """
        Scan the first 20 rows for a header row whose cells match recognisable
        column keywords.  Returns (header_row_1based, col_map).
        """
        for row_idx, row in enumerate(ws.iter_rows(max_row=20, values_only=True), 1):
            mapped: dict[str, int] = {}
            used_cols: set[int] = set()
            for col_idx, cell in enumerate(row):
                if cell is None:
                    continue
                cn = _norm(str(cell))
                for canon, kws in hints.items():
                    if canon in mapped:
                        continue
                    if col_idx in used_cols:
                        continue
                    for kw in kws:
                        if _norm(kw) in cn:
                            mapped[canon] = col_idx
                            used_cols.add(col_idx)
                            break
            if len(mapped) >= 4:
                return row_idx, {**defaults, **mapped}
        log.warning(
            "Column auto-detection failed — using default column indices: %s",
            defaults,
        )
        return 11, defaults

    @staticmethod
    def _post_validate_columns(
        ws,
        col_map: dict,
        data_start_row: int,
    ) -> dict:
        """
        BUG 2 FIX: Validate that the mapped 'quantity' column actually holds
        numeric values in the data rows.  If it doesn't (e.g. it contains unit
        abbreviations like 'шт' or 'kg/кг'), try the 'unit' column instead.

        This handles the manba.xlsx / DBT3004 case where the data values for
        'unit' and 'quantity' are physically swapped relative to the header.
        """
        c_qty  = col_map.get("quantity", 4)
        c_unit = col_map.get("unit", 3)

        numeric_in_qty  = 0
        numeric_in_unit = 0
        checked = 0

        for row in ws.iter_rows(min_row=data_start_row, max_row=data_start_row + 4,
                                 values_only=True):
            if row[0] is None:
                continue
            if _norm(str(row[0])) == "всего":
                continue
            checked += 1
            qty_val  = _safe_get(row, c_qty,  None)
            unit_val = _safe_get(row, c_unit, None)

            if _is_numeric(qty_val):
                numeric_in_qty += 1
            if _is_numeric(unit_val):
                numeric_in_unit += 1

        if checked > 0 and numeric_in_qty == 0 and numeric_in_unit > 0:
            log.warning(
                "BUG2: quantity column %d has no numeric values but unit "
                "column %d does — swapping quantity and unit columns.",
                c_qty, c_unit,
            )
            new_map = dict(col_map)
            new_map["quantity"] = c_unit
            new_map["unit"]     = c_qty
            return new_map

        return col_map

    # ── Header rows ─────────────────────────────────────────────────────────

    @staticmethod
    def _parse_header_rows(ws, data: DeclarationData) -> None:
        """
        Extract contract number, consignee, weights etc. from rows 1–15.

        COMPANY NAME FIX:
          The consignee (and sender) name is passed through _clean_company_name()
          to remove legal-form prefixes (OOO/ООО), surrounding quotes, and
          tax-ID numbers, producing a clean "Name LLC" string.
        """
        for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
            for cell in row:
                if cell is None:
                    continue
                s = str(cell).strip()
                lo = s.lower()
                if lo.startswith("invoice date"):
                    data.invoice_date = s.split(":", 1)[-1].strip()
                elif lo.startswith("invoice №") or lo.startswith("invoice no"):
                    data.contract_no = s.split(":", 1)[-1].strip()
                elif lo.startswith("sender:") or "consigneer:" in lo:
                    if not data.sender_name:
                        part = s.split(":", 1)[-1].strip()
                        # Strip "ОТПРАВИТЕЛЬ / Consigneer:" prefix if present
                        if "/" in part:
                            part = part.split("/", 1)[-1].strip()
                        data.sender_name = _clean_company_name(part)
                elif lo.startswith("consignee:"):
                    raw = s.split(":", 1)[-1].strip()
                    data.consignee = _clean_company_name(raw)
                    log.info(
                        "Consignee raw: %r  →  cleaned: %r",
                        raw, data.consignee,
                    )
                elif lo.startswith("adress:") or lo.startswith("address:"):
                    if not data.consignee_address and data.consignee:
                        data.consignee_address = s.split(":", 1)[-1].strip()
                    elif not data.sender_address:
                        data.sender_address = s.split(":", 1)[-1].strip()
                elif lo.startswith("delivery terms"):
                    data.delivery_terms = s.split(":", 1)[-1].strip()

    # ── Totals row ─────────────────────────────────────────────────────────

    @staticmethod
    def _fill_totals_from_sheet(ws, data: DeclarationData, col_map: dict) -> None:
        """
        Read the 'Всего' totals row for packages, net weight, gross weight.

        Layout B (single-sheet) note:
        The Всего row contains two 'Всего' labels — one for the invoice total
        and one for the packing total.  The packing totals (packages, gross,
        net) follow immediately after the SECOND 'Всего' label in the row.

        Example (manba.xlsx DBT3004):
          ['Всего', 748, 267.95, 'Всего', 21, 290, 243]
                                  ^second → packages=21, gross=290, net=243

        Example (yangi_manba.xlsx ATAF1305):
          ['Всего', 11701.772, 'Всего', 908, 10679, 9771]
                                ^second → packages=908, gross=10679, net=9771
        """
        for row in ws.iter_rows(values_only=True):
            if row[0] is None:
                continue
            if _norm(str(row[0])) != "всего":
                continue

            vsego_indices = [
                i for i, v in enumerate(row)
                if v is not None and _norm(str(v)) == "всего"
            ]

            if len(vsego_indices) >= 2:
                packing_start = vsego_indices[1]
                numeric_vals: list[float] = []
                for ci in range(packing_start + 1, len(row)):
                    v = row[ci]
                    if v is not None and _is_numeric(v):
                        numeric_vals.append(float(v))
                        if len(numeric_vals) == 3:
                            break

                if len(numeric_vals) >= 3:
                    data.packages     = int(round(numeric_vals[0]))
                    data.gross_weight = numeric_vals[1]
                    data.net_weight   = numeric_vals[2]
                else:
                    log.warning(
                        "Could not find 3 numeric values after second 'Всего' "
                        "(found %d): %s", len(numeric_vals), numeric_vals
                    )
                log.info(
                    "Totals (packing start col %d): packages=%d  "
                    "gross=%.1f  net=%.1f",
                    packing_start, data.packages, data.gross_weight, data.net_weight,
                )
            else:
                data.packages     = _int_cell(row,   col_map.get("packages",     4))
                data.net_weight   = _float_cell(row, col_map.get("net_weight",   5))
                data.gross_weight = _float_cell(row, col_map.get("gross_weight", 6))
                log.info(
                    "Totals (col-map): packages=%d  gross=%.1f  net=%.1f",
                    data.packages, data.gross_weight, data.net_weight,
                )
            break

    @staticmethod
    def _parse_totals(ws, col_map: dict) -> dict:
        """Return raw totals dict (kept for potential future use)."""
        return {}

    # ── Items ─────────────────────────────────────────────────────────────

    @staticmethod
    def _parse_items(
        ws,
        data: DeclarationData,
        col_map: dict,
        data_start_row: int,
    ) -> None:
        c_no    = col_map.get("item_no",     0)
        c_desc  = col_map.get("description", 1)
        c_hs    = col_map.get("hs_code",     2)
        c_unit  = col_map.get("unit",        3)
        c_qty   = col_map.get("quantity",    4)
        c_upric = col_map.get("unit_price",  5)
        c_total = col_map.get("total",       6)

        log.info(
            "Item parse — columns: item_no=%d desc=%d hs=%d unit=%d qty=%d "
            "unit_price=%d total=%d  start_row=%d",
            c_no, c_desc, c_hs, c_unit, c_qty, c_upric, c_total, data_start_row,
        )

        for row in ws.iter_rows(min_row=data_start_row, values_only=True):
            if row[0] is None:
                continue
            if _norm(str(row[0])) == "всего":
                break

            try:
                item_no = int(float(str(row[c_no]).strip()))
            except (TypeError, ValueError, IndexError):
                continue

            hs_raw     = str(_safe_get(row, c_hs, "")).strip()
            hs_code    = normalize_hs(hs_raw)
            ru_name    = str(_safe_get(row, c_desc, "")).strip()
            unit_raw   = str(_safe_get(row, c_unit, "шт")).strip()
            quantity   = _int_cell(row, c_qty)
            unit_price = round(_float_cell(row, c_upric), 4)
            total      = round(_float_cell(row, c_total), 2)

            # Derive total from unit_price × quantity if the cell is empty/zero
            if total == 0.0 and unit_price > 0 and quantity > 0:
                total = round(unit_price * quantity, 2)

            # BUG 2 FIX: derive quantity if still 0 but price info is available
            if quantity == 0 and unit_price > 0 and total > 0:
                derived = round(total / unit_price)
                log.warning(
                    "Item %d (HS %s): qty col %d returned 0 — "
                    "derived from total/unit_price: %d (%.2f / %.4f)",
                    item_no, hs_code, c_qty, derived, total, unit_price,
                )
                quantity = derived
            elif quantity == 0:
                log.warning(
                    "Item %d (HS %s): quantity is 0 and cannot be derived "
                    "(unit_price=%.4f, total=%.2f) — check column mapping.",
                    item_no, hs_code, unit_price, total,
                )

            if isinstance(quantity, float):
                quantity = round(quantity)

            zh_name = get_chinese_name(hs_code, ru_name)

            data.items.append(ProductItem(
                item_no      = item_no,
                hs_code      = hs_code,
                russian_name = ru_name,
                chinese_name = zh_name,
                quantity     = int(quantity),
                unit_price   = unit_price,
                total_price  = total,
                unit         = unit_raw,
            ))


# ── Utility functions ─────────────────────────────────────────────────────

def _norm(text: str) -> str:
    """
    Normalise text for column-header matching:
    lowercase + collapse whitespace + normalise apostrophe variants.

    BUG 2: Replace curly apostrophe (U+2019, U+02BC) with ASCII (U+0027)
    so "Q\u2019ty" and "Q'ty" both match the hint "q'ty".
    """
    text = text.replace("\u2019", "'").replace("\u02BC", "'")
    return re.sub(r"\s+", "", text.lower())


def _is_numeric(val) -> bool:
    """Return True if val can be converted to a positive float."""
    if val is None:
        return False
    try:
        return float(str(val).strip()) > 0
    except (ValueError, TypeError):
        return False


def _safe_get(row: tuple, idx: int, default=""):
    try:
        v = row[idx]
        return v if v is not None else default
    except IndexError:
        return default


def _int_cell(row: tuple, idx: int) -> int:
    try:
        val = row[idx]
        if val is None:
            return 0
        return int(float(str(val).strip()))
    except (TypeError, ValueError, IndexError):
        return 0


def _float_cell(row: tuple, idx: int) -> float:
    try:
        val = row[idx]
        if val is None:
            return 0.0
        return float(str(val).strip())
    except (TypeError, ValueError, IndexError):
        return 0.0
